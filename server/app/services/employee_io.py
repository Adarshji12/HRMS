"""Employees to and from Excel: the sample file, the export, and the import check.

One column list drives all three, and it carries every field Add employee
does — basic details, personal, address, statutory, bank, employment and shift
— so a filled sample creates the same complete record as the ten-step form.

The import is checked before anything is written. ``check`` reads every row,
resolves names (department, designation, team, role, shift, manager) against
this company's own records, and validates the HR half with the very schemas the
HR record's own endpoints use, so a PAN or IFSC that the record would refuse is
reported at the check rather than after the person exists. The endpoint then
creates only the ready rows, through the same code paths as adding one person
by hand: seats, the employee-ID series and the audit log apply exactly as they
do to a single hire.

Exports leave PAN, Aadhaar and bank account blank. They are stored encrypted
and shown masked everywhere in the app; a spreadsheet of them in the clear is
exactly the copy the encryption exists to prevent.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import EmailStr, TypeAdapter, ValidationError
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Department, Designation, OrgRole, Team, User
from app.models.attendance import Shift
from app.models.corehr import EmployeeProfile
from app.schemas.corehr import EmployeeProfileIn, EmploymentIn

MAX_ROWS = 1000
MAX_BYTES = 4 * 1024 * 1024

#: The example row on the sample sheet. Left as it is, it is skipped — so a
#: file filled in below it imports without anybody having to delete it first.
EXAMPLE_EMAIL = "rahul.sharma@example.com"

STATUS_LABELS = {"active": "Active", "pending_setup": "Pre-joining", "inactive": "Inactive"}
_STATUS_FROM = {
    "active": "active", "pre joining": "pending_setup", "prejoining": "pending_setup",
    "pending setup": "pending_setup", "inactive": "inactive",
}

#: Fixed choices, shown as drop-downs and read back by label or by value.
CHOICES: dict[str, list[tuple[str, str]]] = {
    "status": [("active", "Active"), ("pending_setup", "Pre-joining")],
    "gender": [("male", "Male"), ("female", "Female"), ("other", "Other"),
               ("undisclosed", "Prefer not to say")],
    "marital_status": [("single", "Single"), ("married", "Married"), ("divorced", "Divorced"),
                       ("widowed", "Widowed"), ("other", "Other")],
    "blood_group": [(b, b) for b in ("A+", "A-", "B+", "B-", "AB+", "AB-", "O+", "O-")],
    "emergency_relationship": [(r, r) for r in ("Father", "Mother", "Spouse", "Sibling", "Son",
                                                "Daughter", "Friend", "Other")],
    "permanent_same": [("yes", "Yes"), ("no", "No")],
    "employment_type": [("full_time", "Full-time"), ("part_time", "Part-time"),
                        ("contract", "Contract"), ("intern", "Intern"),
                        ("consultant", "Consultant")],
    "attendance_type": [("office", "Office"), ("hybrid", "Hybrid"), ("remote", "Remote"),
                        ("field", "Field")],
    "bgv_status": [("pending", "Pending"), ("in_progress", "In progress"),
                   ("cleared", "Cleared"), ("flagged", "Flagged")],
}


@dataclass(frozen=True)
class Column:
    key: str
    header: str
    group: str
    required: bool = False
    width: int = 18
    help: str = ""


def _c(group: str, key: str, header: str, width: int = 18, help: str = "", required: bool = False) -> Column:
    return Column(key=key, header=header, group=group, required=required, width=width, help=help)


COLUMNS: tuple[Column, ...] = (
    # Basic details — what the account needs.
    _c("Basic details", "first_name", "First name", 16, "As on their government ID.", True),
    _c("Basic details", "middle_name", "Middle name", 14),
    _c("Basic details", "last_name", "Last name", 16),
    _c("Basic details", "email", "Work email", 30, "Their sign-in. Must be unique in this company.", True),
    _c("Basic details", "phone", "Mobile", 16),
    _c("Basic details", "employee_code", "Employee ID", 14, "Blank issues the next ID from HR Settings."),
    _c("Basic details", "joined_on", "Date of joining", 16, "A date, or YYYY-MM-DD / DD-MM-YYYY. Blank means today."),
    _c("Basic details", "status", "Status", 14, "Active or Pre-joining. Blank means Active."),
    _c("Basic details", "timezone", "Time zone", 16, "For example Asia/Kolkata. Blank means Asia/Kolkata."),
    # Personal.
    _c("Personal", "date_of_birth", "Date of birth", 15, "A date, or YYYY-MM-DD / DD-MM-YYYY."),
    _c("Personal", "gender", "Gender", 14),
    _c("Personal", "marital_status", "Marital status", 14),
    _c("Personal", "blood_group", "Blood group", 12),
    _c("Personal", "nationality", "Nationality", 14),
    _c("Personal", "father_name", "Father's / spouse's name", 22, "PF nomination and Form 11 ask for it."),
    _c("Personal", "personal_email", "Personal email", 26, "Kept apart from work email; it outlives employment."),
    _c("Personal", "alternate_phone", "Alternate phone", 16),
    _c("Personal", "emergency_name", "Emergency contact name", 22),
    _c("Personal", "emergency_relationship", "Emergency relationship", 18),
    _c("Personal", "emergency_phone", "Emergency phone", 16),
    # Address.
    _c("Address", "current_address", "Current address", 32),
    _c("Address", "current_city", "Current city", 16),
    _c("Address", "current_state", "Current state", 16),
    _c("Address", "current_postal_code", "Current PIN code", 14),
    _c("Address", "current_country", "Current country", 14, "Two-letter code. Blank means IN."),
    _c("Address", "permanent_same", "Permanent same as current", 22, "Yes copies the current address."),
    _c("Address", "permanent_address", "Permanent address", 32),
    _c("Address", "permanent_city", "Permanent city", 16),
    _c("Address", "permanent_state", "Permanent state", 16),
    _c("Address", "permanent_postal_code", "Permanent PIN code", 16),
    _c("Address", "permanent_country", "Permanent country", 16, "Two-letter code. Blank means IN."),
    # Statutory.
    _c("Statutory", "pan", "PAN", 14, "10 characters, e.g. ABCDE1234F. Stored encrypted."),
    _c("Statutory", "aadhaar", "Aadhaar", 16, "12 digits; spaces are fine. Stored encrypted."),
    _c("Statutory", "uan", "UAN", 16),
    _c("Statutory", "pf_number", "PF number", 18),
    _c("Statutory", "esi_number", "ESI number", 16),
    # Bank.
    _c("Bank", "bank_name", "Bank name", 18),
    _c("Bank", "bank_branch", "Bank branch", 16),
    _c("Bank", "bank_ifsc", "IFSC", 14, "11 characters, e.g. HDFC0001234."),
    _c("Bank", "bank_account", "Account number", 20, "Stored encrypted. Format the cell as Text so Excel keeps leading zeros."),
    _c("Bank", "bank_account_name", "Account holder name", 22, "Blank means the employee's name."),
    # Employment & placement.
    _c("Employment", "role", "Role", 18, "One of this company's roles. Blank means Employee."),
    _c("Employment", "department", "Department", 20, "Must match a Department master by name."),
    _c("Employment", "designation", "Designation", 20, "Must match a Designation master by name."),
    _c("Employment", "team", "Team", 18, "Must match a team by name."),
    _c("Employment", "manager_email", "Manager email", 28, "An existing employee, or someone else in this file."),
    _c("Employment", "employment_type", "Employment type", 16),
    _c("Employment", "work_location", "Work location", 16),
    _c("Employment", "grade", "Grade / band", 12),
    _c("Employment", "probation_months", "Probation (months)", 16, "0 if they join confirmed; up to 36."),
    _c("Employment", "notice_period_days", "Notice period (days)", 18, "Up to 365."),
    _c("Employment", "bgv_status", "Background verification", 20),
    # Attendance.
    _c("Attendance", "attendance_type", "Attendance type", 16, "Office, Hybrid, Remote or Field."),
    _c("Attendance", "shift", "Shift", 18, "Must match a shift by name. Starts on the joining date."),
)

GROUPS = list(dict.fromkeys(c.group for c in COLUMNS))
_GROUP_FILLS = {
    "Basic details": "0F3D66", "Personal": "275D8C", "Address": "2F6F73", "Statutory": "6B4E9A",
    "Bank": "5B6B2F", "Employment": "8A5A1F", "Attendance": "4A5568",
}
_PROFILE_KEYS = (
    "date_of_birth", "gender", "marital_status", "blood_group", "nationality", "father_name",
    "personal_email", "alternate_phone", "emergency_name", "emergency_relationship",
    "emergency_phone", "current_address", "current_city", "current_state", "current_postal_code",
    "current_country", "permanent_address", "permanent_city", "permanent_state",
    "permanent_postal_code", "permanent_country", "pan", "aadhaar", "uan", "pf_number",
    "esi_number", "bank_name", "bank_branch", "bank_ifsc", "bank_account", "bank_account_name",
)
_EMPLOYMENT_KEYS = ("employment_type", "work_location", "grade", "attendance_type",
                    "probation_months", "notice_period_days", "bgv_status")
#: Never written to an export: stored encrypted, shown masked.
_SECRET = {"pan", "aadhaar", "bank_account"}
_DATE_KEYS = {"joined_on", "date_of_birth"}
_TEXT_KEYS = {"phone", "alternate_phone", "emergency_phone", "aadhaar", "bank_account", "uan",
              "pf_number", "esi_number", "current_postal_code", "permanent_postal_code",
              "employee_code"}

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_EXAMPLE_FONT = Font(italic=True, color="7A8794")


def _norm(text: object) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(text or "").lower()).strip()


_HEADER_KEYS = {_norm(c.header): c.key for c in COLUMNS} | {_norm(c.key): c.key for c in COLUMNS}
def _ckey(text: object) -> str:
    """Like _norm, but keeps + and - and drops spaces: "O+" and "O-" are
    different blood groups, and "B +" is still B+."""
    return re.sub(r"[^a-z0-9+\-]+", "", str(text or "").lower())


def _choice_index(pairs: list[tuple[str, str]]) -> dict[str, str]:
    """Typed text to stored value. The exact form first; the looser _norm form
    ("pre joining" for "Pre-joining") only where it points at one value."""
    exact = {_ckey(v): v for v, _ in pairs} | {_ckey(label): v for v, label in pairs}
    loose: dict[str, set[str]] = {}
    for v, label in pairs:
        for k in (_norm(v), _norm(label)):
            loose.setdefault(k, set()).add(v)
    return {k: next(iter(vs)) for k, vs in loose.items() if len(vs) == 1} | exact


_CHOICE_FROM = {key: _choice_index(pairs) for key, pairs in CHOICES.items()}

#: Checked exactly as the create-employee endpoint checks it, so the import never
#: passes an address the API would then refuse, or refuses one it would take.
_EMAIL = TypeAdapter(EmailStr)

#: The four built-in levels by the names people type, for a company that has
#: not yet opened the Roles screen (where named roles are first created).
_LEVELS = {"employee": "employee", "standard": "employee", "staff": "employee",
           "manager": "manager", "administrator": "admin", "admin": "admin"}


@dataclass
class Lookups:
    departments: dict[str, Department]
    designations: dict[str, Designation]
    teams: dict[str, Team]
    roles: dict[str, OrgRole]
    shifts: dict[str, Shift]
    users_by_email: dict[str, User]
    codes: set[str]


def lookups(db: Session, org_id) -> Lookups:
    def by_name(model):
        rows = db.scalars(select(model).where(model.org_id == org_id)).all()
        return {_norm(r.name): r for r in rows}

    roles = {
        _norm(r.name): r
        for r in db.scalars(select(OrgRole).where(OrgRole.org_id == org_id)).all()
        if r.base_role != "owner"
    }
    users = db.scalars(select(User).where(User.org_id == org_id)).all()
    return Lookups(
        departments=by_name(Department),
        designations=by_name(Designation),
        teams=by_name(Team),
        roles=roles,
        shifts=by_name(Shift),
        users_by_email={u.email.lower(): u for u in users},
        codes={u.employee_code.lower() for u in users if u.employee_code},
    )


# --------------------------------------------------------------------- write --


def _col(key: str) -> int:
    return next(i for i, c in enumerate(COLUMNS, start=1) if c.key == key)


def _sheet_header(ws) -> None:
    """Row 1: the section bands. Row 2: the column headers. Data from row 3."""
    start = 1
    for group in GROUPS:
        span = sum(1 for c in COLUMNS if c.group == group)
        fill = PatternFill("solid", fgColor=_GROUP_FILLS.get(group, "0F3D66"))
        ws.cell(row=1, column=start, value=group.upper())
        if span > 1:
            ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=start + span - 1)
        for i in range(start, start + span):
            ws.cell(row=1, column=i).fill = fill
            head = ws.cell(row=2, column=i)
            head.fill = fill
        ws.cell(row=1, column=start).font = _HEADER_FONT
        ws.cell(row=1, column=start).alignment = Alignment(horizontal="left", vertical="center", indent=1)
        start += span
    for i, col in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=2, column=i, value=f"{col.header} *" if col.required else col.header)
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = col.width
    ws.row_dimensions[1].height = 20
    ws.row_dimensions[2].height = 30
    ws.freeze_panes = "C3"


def _to_bytes(wb: Workbook) -> bytes:
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _format_cells(ws, first_row: int, last_row: int) -> None:
    for key in _DATE_KEYS:
        letter = get_column_letter(_col(key))
        for r in range(first_row, last_row + 1):
            ws[f"{letter}{r}"].number_format = "yyyy-mm-dd"
    # Text, so Excel keeps leading zeros on accounts and PINs and does not turn
    # a 12-digit Aadhaar into 1.23457E+11.
    for key in _TEXT_KEYS:
        letter = get_column_letter(_col(key))
        for r in range(first_row, last_row + 1):
            ws[f"{letter}{r}"].number_format = "@"


def template(lk: Lookups) -> bytes:
    """The sample file: section bands, headers, one example row, drop-downs and instructions."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Employees"
    _sheet_header(ws)
    _format_cells(ws, 3, MAX_ROWS + 2)

    first = lambda table: next(iter(table.values())).name if table else ""  # noqa: E731
    example = {
        "first_name": "Rahul", "last_name": "Sharma", "email": EXAMPLE_EMAIL,
        "phone": "9876543210", "joined_on": date.today(), "status": "Active",
        "timezone": "Asia/Kolkata", "date_of_birth": date(1995, 6, 15), "gender": "Male",
        "marital_status": "Single", "blood_group": "B+", "nationality": "Indian",
        "father_name": "Suresh Sharma", "personal_email": "rahul.personal@example.com",
        "emergency_name": "Suresh Sharma", "emergency_relationship": "Father",
        "emergency_phone": "9876500000", "current_address": "12 MG Road",
        "current_city": "Pune", "current_state": "Maharashtra", "current_postal_code": "411001",
        "current_country": "IN", "permanent_same": "Yes", "pan": "ABCDE1234F",
        "uan": "100200300400", "bank_name": "HDFC Bank", "bank_branch": "MG Road",
        "bank_ifsc": "HDFC0001234", "bank_account": "50100012345678",
        "role": "Employee", "department": first(lk.departments), "designation": first(lk.designations),
        "employment_type": "Full-time", "work_location": "Pune office", "grade": "L2",
        "probation_months": 6, "notice_period_days": 30, "bgv_status": "Pending",
        "attendance_type": "Office", "shift": first(lk.shifts),
    }
    for i, col in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=3, column=i, value=example.get(col.key) or None)
        cell.font = _EXAMPLE_FONT

    # Drop-downs, fed from a hidden sheet so long master lists fit.
    lists = wb.create_sheet("Lists")
    lists.sheet_state = "hidden"
    options: dict[str, list[str]] = {k: [label for _, label in v] for k, v in CHOICES.items()}
    options |= {
        "role": sorted(r.name for r in lk.roles.values()) or ["Employee", "Manager", "Administrator"],
        "department": sorted(d.name for d in lk.departments.values()),
        "designation": sorted(d.name for d in lk.designations.values()),
        "team": sorted(t.name for t in lk.teams.values()),
        "shift": sorted(s.name for s in lk.shifts.values()),
    }
    for n, (key, values) in enumerate(options.items(), start=1):
        if not values:
            continue
        letter = get_column_letter(n)
        lists.cell(row=1, column=n, value=key)
        for r, value in enumerate(values, start=2):
            lists.cell(row=r, column=n, value=value)
        dv = DataValidation(type="list", formula1=f"=Lists!${letter}$2:${letter}${len(values) + 1}",
                            allow_blank=True, showErrorMessage=False)
        target = get_column_letter(_col(key))
        dv.add(f"{target}3:{target}{MAX_ROWS + 2}")
        ws.add_data_validation(dv)

    guide = wb.create_sheet("Instructions")
    guide.column_dimensions["A"].width = 16
    guide.column_dimensions["B"].width = 28
    guide.column_dimensions["C"].width = 10
    guide.column_dimensions["D"].width = 70
    guide.append(["How to fill this file"])
    guide["A1"].font = Font(bold=True, size=13)
    guide.append([])
    for line in (
        "One employee per row on the Employees sheet, from row 3. Columns marked * are required; "
        "everything else is optional and can be filled in later on the employee's record.",
        f"Row 3 is an example ({EXAMPLE_EMAIL}). It is skipped if left as it is — type over it or below it.",
        "Columns with a drop-down only accept those values. Department, Designation, Team, Role and "
        "Shift must already be set up in the app.",
        "Nothing is saved until you check the file and press Import — rows with problems are listed "
        "with the reason and left out.",
        "Salary structure, assets and documents are not in this file; add them on the employee's "
        "record after the import.",
        f"Up to {MAX_ROWS} employees per file.",
    ):
        guide.append([line])
        guide.merge_cells(start_row=guide.max_row, start_column=1, end_row=guide.max_row, end_column=4)
    guide.append([])
    guide.append(["Section", "Column", "Required", "What to enter"])
    for c in (guide.cell(row=guide.max_row, column=i) for i in range(1, 5)):
        c.font = Font(bold=True)
    for col in COLUMNS:
        choices = CHOICES.get(col.key)
        what = col.help
        if choices and not what:
            what = "One of: " + ", ".join(label for _, label in choices) + "."
        guide.append([col.group, col.header, "Yes" if col.required else "", what])
    for row in guide.iter_rows(min_row=3):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    return _to_bytes(wb)


def export(people: list[User], lk: Lookups, managers: dict, profiles: dict) -> bytes:
    """Every employee given, in the import's own columns. PAN, Aadhaar and bank
    account are left blank — see the module docstring."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Employees"
    _sheet_header(ws)
    _format_cells(ws, 3, len(people) + 2)
    dept = {d.id: d.name for d in lk.departments.values()}
    desig = {d.id: d.name for d in lk.designations.values()}
    team = {t.id: t.name for t in lk.teams.values()}
    role = {r.id: r.name for r in lk.roles.values()}
    label = {k: dict(v) for k, v in CHOICES.items()}
    for r, u in enumerate(people, start=3):
        first, middle, last = split_name(u.full_name)
        manager = managers.get(u.manager_id)
        profile = profiles.get(u.id)
        values: dict[str, object] = {
            "first_name": first, "middle_name": middle, "last_name": last, "email": u.email,
            "phone": u.phone, "employee_code": u.employee_code, "joined_on": u.joined_on,
            "status": STATUS_LABELS.get(u.status or "active", u.status), "timezone": u.timezone,
            "role": role.get(u.role_id) or (u.role or "").capitalize(),
            "department": dept.get(u.department_id), "designation": desig.get(u.designation_id),
            "team": team.get(u.team_id), "manager_email": manager.email if manager else None,
        }
        if profile is not None:
            for key in (*_PROFILE_KEYS, *_EMPLOYMENT_KEYS):
                if key in _SECRET:
                    continue
                value = getattr(profile, key, None)
                values[key] = label.get(key, {}).get(value, value) if isinstance(value, str) else value
        for i, col in enumerate(COLUMNS, start=1):
            value = values.get(col.key)
            ws.cell(row=r, column=i, value=value if value not in ("", None) else None)
    ws.auto_filter.ref = f"A2:{get_column_letter(len(COLUMNS))}{max(2, len(people) + 2)}"
    return _to_bytes(wb)


def split_name(full: str) -> tuple[str, str, str]:
    parts = (full or "").split()
    if len(parts) <= 1:
        return (parts[0] if parts else "", "", "")
    if len(parts) == 2:
        return parts[0], "", parts[1]
    return parts[0], " ".join(parts[1:-1]), parts[-1]


# ---------------------------------------------------------------------- read --


class FileRefused(ValueError):
    """The upload cannot be read at all — as opposed to a row that is wrong."""


def read_rows(content: bytes, filename: str) -> list[tuple[int, dict[str, object]]]:
    """(sheet row number, {column key: raw value}) for every non-empty row.

    The header row is found rather than assumed: the sample has a band of
    section names above it, and a hand-made file may not.
    """
    if len(content) > MAX_BYTES:
        raise FileRefused(f"The file is larger than {MAX_BYTES // (1024 * 1024)} MB.")
    name = (filename or "").lower()
    if name.endswith(".csv"):
        text = content.decode("utf-8-sig", errors="replace")
        table = list(csv.reader(io.StringIO(text)))
    elif name.endswith((".xlsx", ".xlsm")):
        try:
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:  # openpyxl raises a zoo of types for a bad file
            raise FileRefused("That file could not be opened as an Excel workbook.") from exc
        ws = wb["Employees"] if "Employees" in wb.sheetnames else wb.worksheets[0]
        table = [list(r) for r in ws.iter_rows(values_only=True)]
    else:
        raise FileRefused("Upload the sample .xlsx file (or a .csv with the same columns).")

    if not table:
        raise FileRefused("The file is empty.")
    header_at = next(
        (i for i, r in enumerate(table[:5]) if "email" in [_HEADER_KEYS.get(_norm(h)) for h in r]),
        None,
    )
    if header_at is None:
        raise FileRefused(
            "Could not find the column headers from the sample file — the row with "
            "First name and Work email."
        )
    keys = [_HEADER_KEYS.get(_norm(h)) for h in table[header_at]]
    missing = [c.header for c in COLUMNS if c.required and c.key not in keys]
    if missing:
        raise FileRefused("The header row is missing: " + ", ".join(missing) + ".")

    rows: list[tuple[int, dict[str, object]]] = []
    for number, raw in enumerate(table[header_at + 1:], start=header_at + 2):
        values = {k: v for k, v in zip(keys, raw) if k}
        if all(v is None or str(v).strip() == "" for v in values.values()):
            continue
        rows.append((number, values))
    if len(rows) > MAX_ROWS:
        raise FileRefused(f"The file has {len(rows)} employees; the limit is {MAX_ROWS} per file.")
    return rows


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)  # a phone number Excel stored as 9876543210.0
    if isinstance(value, (datetime, date)):
        return value.isoformat()[:10]
    return str(value).strip()


def _date(value: object) -> date | None:
    if value is None or _text(value) == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value)
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%Y/%m/%d", "%d.%m.%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"“{text}” is not a date — use YYYY-MM-DD")


_HEADER_BY_KEY = {c.key: c.header for c in COLUMNS}


def _schema_errors(exc: ValidationError) -> list[str]:
    """Pydantic's errors, named by the sheet's own column headers."""
    out = []
    for err in exc.errors():
        key = str(err["loc"][0]) if err.get("loc") else ""
        message = err.get("msg", "is not valid").removeprefix("Value error, ")
        out.append(f"{_HEADER_BY_KEY.get(key, key)}: {message}.")
    return out


@dataclass
class Checked:
    row: int
    name: str
    email: str
    errors: list[str] = field(default_factory=list)
    skipped: str | None = None
    #: Ready-to-create account fields, for the rows with no errors.
    values: dict = field(default_factory=dict)
    profile: EmployeeProfileIn | None = None
    employment: EmploymentIn | None = None
    probation_months: int = 0
    shift_id: object = None
    manager_email: str | None = None
    #: Set at import when the account was created but a later step failed.
    warnings: list[str] = field(default_factory=list)

    @property
    def state(self) -> str:
        return "skipped" if self.skipped else ("error" if self.errors else "ready")


def check(rows: list[tuple[int, dict]], lk: Lookups, *, can_people: bool = True) -> list[Checked]:
    """Every row, resolved and validated. Nothing is written."""
    seen_emails: dict[str, int] = {}
    seen_codes: dict[str, int] = {}
    in_file = {_text(v.get("email")).lower() for _, v in rows if _text(v.get("email"))}
    out: list[Checked] = []

    for number, v in rows:
        first = _text(v.get("first_name"))
        middle = _text(v.get("middle_name"))
        last = _text(v.get("last_name"))
        name = " ".join(p for p in (first, middle, last) if p)
        email = _text(v.get("email")).lower()
        item = Checked(row=number, name=name, email=email)
        out.append(item)

        if email == EXAMPLE_EMAIL:
            item.skipped = "The sample row — left as it was."
            continue

        if not first:
            item.errors.append("First name is required.")
        if not email:
            item.errors.append("Work email is required.")
        else:
            try:
                _EMAIL.validate_python(email)
            except ValidationError:
                item.errors.append(f"“{email}” is not a valid email address.")
            if email in lk.users_by_email:
                item.errors.append("An employee with this email already exists.")
            if email in seen_emails:
                item.errors.append(f"The same email is on row {seen_emails[email]} too.")
            seen_emails.setdefault(email, number)

        code = _text(v.get("employee_code"))
        if code:
            if code.lower() in lk.codes:
                item.errors.append(f"Employee ID {code} is already in use.")
            if code.lower() in seen_codes:
                item.errors.append(f"Employee ID {code} is on row {seen_codes[code.lower()]} too.")
            seen_codes.setdefault(code.lower(), number)

        def when(key: str) -> date | None:
            try:
                return _date(v.get(key))
            except ValueError as exc:
                item.errors.append(f"{_HEADER_BY_KEY[key]}: {exc}.")
                return None

        joined = when("joined_on") or date.today()

        def choice(key: str) -> str | None:
            text = _text(v.get(key))
            if not text:
                return None
            found = _CHOICE_FROM[key].get(_ckey(text)) or _CHOICE_FROM[key].get(_norm(text))
            if found is None:
                labels = ", ".join(label for _, label in CHOICES[key])
                item.errors.append(f"{_HEADER_BY_KEY[key]} “{text}” — use one of: {labels}.")
            return found

        status = choice("status") or "active"

        def pick(key: str, table: dict, what: str):
            text = _text(v.get(key))
            if not text:
                return None
            found = table.get(_norm(text))
            if found is None:
                item.errors.append(f"{what} “{text}” is not set up — add it in the app first.")
            return found

        department = pick("department", lk.departments, "Department")
        designation = pick("designation", lk.designations, "Designation")
        team = pick("team", lk.teams, "Team")
        shift = pick("shift", lk.shifts, "Shift")

        role_text = _text(v.get("role"))
        role = lk.roles.get(_norm(role_text)) if role_text else lk.roles.get("employee")
        level = role.base_role if role else _LEVELS.get(_norm(role_text) or "employee")
        if role_text and role is None and level is None:
            if _norm(role_text) == "owner":
                item.errors.append("Nobody can be imported as the Owner.")
            else:
                item.errors.append(f"Role “{role_text}” is not one of this company's roles.")

        manager = _text(v.get("manager_email")).lower()
        if manager:
            if manager == email:
                item.errors.append("An employee cannot be their own manager.")
            elif manager not in lk.users_by_email and manager not in in_file:
                item.errors.append(f"Manager {manager} is not an employee and is not in this file.")
            item.manager_email = manager

        # The HR record: built and validated with the record's own schemas.
        profile: dict[str, object] = {}
        for key in _PROFILE_KEYS:
            if key == "date_of_birth":
                dob = when("date_of_birth")
                if dob:
                    profile[key] = dob
                continue
            if key in CHOICES:
                found = choice(key)
                if found:
                    profile[key] = found
                continue
            text = _text(v.get(key))
            if text:
                profile[key] = text
        for side in ("current", "permanent"):
            if profile.get(f"{side}_address") or profile.get(f"{side}_city"):
                profile[f"{side}_country"] = str(profile.get(f"{side}_country") or "IN").upper()
        if (choice("permanent_same") or "") == "yes":
            for part in ("address", "city", "state", "postal_code", "country"):
                if profile.get(f"current_{part}"):
                    profile[f"permanent_{part}"] = profile[f"current_{part}"]
        if profile.get("bank_account") and not profile.get("bank_account_name"):
            profile["bank_account_name"] = name

        employment: dict[str, object] = {}
        for key in _EMPLOYMENT_KEYS:
            if key in CHOICES:
                found = choice(key)
                if found:
                    employment[key] = found
                continue
            text = _text(v.get(key))
            if not text:
                continue
            if key in ("probation_months", "notice_period_days"):
                try:
                    employment[key] = int(float(text))
                except ValueError:
                    item.errors.append(f"{_HEADER_BY_KEY[key]}: “{text}” is not a number.")
                    continue
            else:
                employment[key] = text
        months = int(employment.get("probation_months") or 0)
        if months > 0:
            employment["probation_end_on"] = _add_months(joined, months)

        if (profile or employment) and not can_people:
            item.errors.append(
                "Personal, statutory, bank and employment columns need the "
                "“See and edit HR records” permission. Clear them, or ask someone who has it."
            )
        try:
            item.profile = EmployeeProfileIn(**profile) if profile else None
        except ValidationError as exc:
            item.errors.extend(_schema_errors(exc))
        try:
            item.employment = EmploymentIn(**employment) if employment else None
        except ValidationError as exc:
            item.errors.extend(_schema_errors(exc))

        item.probation_months = months
        item.shift_id = shift.id if shift else None
        item.values = {
            "email": email,
            "full_name": name,
            "phone": _text(v.get("phone")) or None,
            "employee_code": code or None,
            "joined_on": joined,
            "status": status,
            "role": level or "employee",
            "role_id": role.id if role else None,
            "department_id": department.id if department else None,
            "designation_id": designation.id if designation else None,
            "team_id": team.id if team else None,
            "timezone": _text(v.get("timezone")) or "Asia/Kolkata",
        }
    return out


def _add_months(start: date, months: int) -> date:
    month = start.month - 1 + months
    year = start.year + month // 12
    month = month % 12 + 1
    days = [31, 29 if year % 4 == 0 and (year % 100 or year % 400 == 0) else 28,
            31, 30, 31, 30, 31, 31, 30, 31, 30, 31][month - 1]
    return date(year, month, min(start.day, days))


def profiles_for(db: Session, user_ids: list) -> dict:
    if not user_ids:
        return {}
    rows = db.scalars(select(EmployeeProfile).where(EmployeeProfile.user_id.in_(user_ids))).all()
    return {p.user_id: p for p in rows}
