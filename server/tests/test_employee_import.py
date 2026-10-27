"""Employees to and from Excel: the sample file, the check, the import, the export."""

from __future__ import annotations

import io
from datetime import date

from openpyxl import Workbook, load_workbook

from app.services.employee_io import COLUMNS, EXAMPLE_EMAIL

from tests.test_employee_hire import _grant_seats

HEADERS = [f"{c.header} *" if c.required else c.header for c in COLUMNS]
KEYS = [c.key for c in COLUMNS]


def _xlsx(rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Employees"
    ws.append(HEADERS)
    for row in rows:
        ws.append([row.get(k) for k in KEYS])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _post(client, tenant, content: bytes, commit: bool, name: str = "people.xlsx"):
    return client.post(
        f"/api/v1/admin/employees/import?commit={'true' if commit else 'false'}",
        headers=tenant.headers(),
        files={"file": (name, content, "application/octet-stream")},
    )


def test_the_sample_file_has_every_column_and_the_example_row(client, tenants):
    acme, _ = tenants
    client.post("/api/v1/org/departments", json={"name": "Sales"}, headers=acme.headers())
    response = client.get("/api/v1/admin/employees/template.xlsx", headers=acme.headers())
    assert response.status_code == 200
    wb = load_workbook(io.BytesIO(response.content))
    assert {"Employees", "Instructions"} <= set(wb.sheetnames)
    ws = wb["Employees"]
    assert [c.value for c in ws[2]] == HEADERS
    assert ws.cell(row=3, column=KEYS.index("email") + 1).value == EXAMPLE_EMAIL
    # Every section of Add employee is in the file.
    assert {"pan", "bank_ifsc", "date_of_birth", "employment_type", "shift"} <= set(KEYS)


def test_the_sample_file_imports_untouched_with_its_example_skipped(client, tenants):
    acme, _ = tenants
    sample = client.get("/api/v1/admin/employees/template.xlsx", headers=acme.headers()).content
    body = _post(client, acme, sample, commit=False).json()
    assert body["skipped"] == 1 and body["ready"] == 0 and body["errors"] == 0


def test_check_writes_nothing_and_commit_creates_the_ready_rows(client, tenants):
    acme, _ = tenants
    _grant_seats(acme.id)
    client.post("/api/v1/org/departments", json={"name": "Sales"}, headers=acme.headers())
    content = _xlsx([
        {"first_name": "Asha", "last_name": "Rao", "email": "asha@acmecorp.com", "department": "sales",
         "joined_on": date(2026, 4, 1), "status": "Pre-joining", "phone": 9876543210},
        {"first_name": "Bad", "email": "not-an-email"},
        {"first_name": "Ravi", "email": "ravi@acmecorp.com", "department": "Marketing"},
    ])

    preview = _post(client, acme, content, commit=False).json()
    assert [r["state"] for r in preview["rows"]] == ["ready", "error", "error"]
    assert "not a valid email" in preview["rows"][1]["messages"][0]
    assert "Marketing" in preview["rows"][2]["messages"][0]
    emails = [u["email"] for u in client.get("/api/v1/admin/employees", headers=acme.headers()).json()]
    assert "asha@acmecorp.com" not in emails

    done = _post(client, acme, content, commit=True).json()
    assert done["created"] == 1
    people = {u["email"]: u for u in client.get("/api/v1/admin/employees", headers=acme.headers()).json()}
    asha = people["asha@acmecorp.com"]
    assert asha["full_name"] == "Asha Rao"
    assert asha["department_name"] == "Sales"
    assert asha["joined_on"] == "2026-04-01"
    assert asha["status"] == "pending_setup"
    assert asha["phone"] == "9876543210"
    assert asha["employee_code"]  # minted from the series
    assert "ravi@acmecorp.com" not in people


def test_a_manager_can_be_someone_else_in_the_same_file(client, tenants):
    acme, _ = tenants
    _grant_seats(acme.id)
    content = _xlsx([
        {"first_name": "Report", "email": "report@acmecorp.com", "manager_email": "boss@acmecorp.com"},
        {"first_name": "Boss", "email": "boss@acmecorp.com"},
    ])
    assert _post(client, acme, content, commit=True).json()["created"] == 2
    people = {u["email"]: u for u in client.get("/api/v1/admin/employees", headers=acme.headers()).json()}
    assert people["report@acmecorp.com"]["manager_id"] == people["boss@acmecorp.com"]["id"]


def test_duplicates_are_caught_in_the_file_and_against_the_company(client, tenants):
    acme, _ = tenants
    content = _xlsx([
        {"first_name": "Dup", "email": "dup@acmecorp.com"},
        {"first_name": "Dup again", "email": "dup@acmecorp.com"},
        {"first_name": "Owner twin", "email": acme.owner.email},
    ])
    rows = _post(client, acme, content, commit=False).json()["rows"]
    assert rows[0]["state"] == "ready"
    assert "row 2" in rows[1]["messages"][0]
    assert any("already exists" in m for m in rows[2]["messages"])


def test_a_file_without_the_headers_is_refused_whole(client, tenants):
    acme, _ = tenants
    wb = Workbook()
    wb.active.append(["Name", "Phone"])
    out = io.BytesIO()
    wb.save(out)
    response = _post(client, acme, out.getvalue(), commit=False)
    assert response.status_code == 422
    assert "headers" in response.json()["detail"]


def test_csv_with_the_same_columns_is_accepted(client, tenants):
    acme, _ = tenants
    _grant_seats(acme.id)
    text = ",".join(HEADERS) + "\nMeena,,Iyer,meena@acmecorp.com" + "," * (len(HEADERS) - 4) + "\n"
    body = _post(client, acme, text.encode(), commit=True, name="people.csv").json()
    assert body["created"] == 1


def test_export_reimports_as_duplicates_not_as_garbage(client, tenants):
    acme, _ = tenants
    exported = client.get("/api/v1/admin/employees/export.xlsx", headers=acme.headers())
    assert exported.status_code == 200
    rows = _post(client, acme, exported.content, commit=False).json()["rows"]
    assert rows and all(any("already exists" in m for m in r["messages"]) for r in rows)


def test_only_administrators_can_import(client, tenants):
    acme, _ = tenants
    response = client.post(
        "/api/v1/admin/employees/import",
        headers=acme.headers(acme.employee_token),
        files={"file": ("p.xlsx", _xlsx([]), "application/octet-stream")},
    )
    assert response.status_code == 403


def test_without_seats_each_row_says_why_it_was_not_created(client, tenants):
    acme, _ = tenants
    content = _xlsx([{"first_name": "Noseat", "email": "noseat@acmecorp.com"}])
    body = _post(client, acme, content, commit=True).json()
    assert body["created"] == 0
    assert body["rows"][0]["state"] == "error"
    assert "subscription" in body["rows"][0]["messages"][0].lower()


def test_the_full_record_is_saved_not_only_the_account(client, tenants):
    acme, _ = tenants
    _grant_seats(acme.id)
    content = _xlsx([{
        "first_name": "Neha", "last_name": "Kulkarni", "email": "neha.k@acmecorp.com",
        "joined_on": date(2026, 4, 1),
        "date_of_birth": "15-06-1995", "gender": "Female", "marital_status": "Married",
        "blood_group": "O+", "nationality": "Indian", "father_name": "Ravi Kulkarni",
        "personal_email": "neha.personal@gmail.com", "emergency_name": "Ravi Kulkarni",
        "emergency_relationship": "Father", "emergency_phone": "9800011111",
        "current_address": "12 MG Road", "current_city": "Pune", "current_state": "Maharashtra",
        "current_postal_code": "411001", "permanent_same": "Yes",
        "pan": "abcde1234f", "aadhaar": "1234 5678 9012", "uan": "100200300400",
        "bank_name": "HDFC Bank", "bank_ifsc": "HDFC0001234", "bank_account": "50100012345678",
        "employment_type": "Full-time", "work_location": "Pune", "grade": "L2",
        "probation_months": 6, "notice_period_days": 30, "bgv_status": "Pending",
        "attendance_type": "Hybrid",
    }])
    body = _post(client, acme, content, commit=True).json()
    assert body["created"] == 1, body
    assert body["rows"][0]["messages"] == []
    people = {u["email"]: u for u in client.get("/api/v1/admin/employees", headers=acme.headers()).json()}
    profile = client.get(f"/api/v1/people/{people['neha.k@acmecorp.com']['id']}", headers=acme.headers()).json()
    assert profile["date_of_birth"] == "1995-06-15"
    assert profile["gender"] == "female"
    assert profile["blood_group"] == "O+"
    assert profile["permanent_city"] == "Pune"          # "same as current" copied
    assert profile["current_country"] == "IN"           # defaulted
    assert profile["has_pan"] and profile["pan_masked"].endswith("234F")
    assert profile["has_bank_account"]
    assert profile["bank_account_name"] == "Neha Kulkarni"  # defaulted to the employee
    assert profile["employment_type"] == "full_time"
    assert profile["attendance_type"] == "hybrid"
    assert profile["notice_period_days"] == 30
    assert profile["probation_end_on"] == "2026-10-01"


def test_hr_fields_are_validated_at_the_check(client, tenants):
    acme, _ = tenants
    content = _xlsx([{
        "first_name": "Bad", "email": "bad.hr@acmecorp.com",
        "pan": "12345", "bank_ifsc": "nope", "gender": "Robot", "probation_months": 99,
    }])
    row = _post(client, acme, content, commit=False).json()["rows"][0]
    assert row["state"] == "error"
    text = " ".join(row["messages"])
    assert "Gender" in text and "PAN" in text and "IFSC" in text and "Probation" in text
